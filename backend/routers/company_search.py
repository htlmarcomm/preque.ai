from fastapi import APIRouter, Depends, Query, UploadFile, File, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import or_
from typing import Optional, List
import openpyxl, io, re
from models.database import get_db, CompanyField, FinancialRecord, ProjectReference

router = APIRouter()

@router.get("/search")
def search_company_data(
    q: Optional[str] = None,
    type: Optional[str] = None,  # "field" | "financial" | "project"
    category: Optional[str] = None,
    year: Optional[str] = None,
    db: Session = Depends(get_db)
):
    results = {"fields": [], "financials": [], "projects": []}
    
    # 1. Search CompanyFields
    if not type or type == "field":
        query = db.query(CompanyField)
        if category:
            query = query.filter(CompanyField.category == category)
        if q:
            search_str = f"%{q}%"
            query = query.filter(
                or_(
                    CompanyField.field_label.ilike(search_str),
                    CompanyField.value.ilike(search_str),
                    CompanyField.aliases.ilike(search_str)
                )
            )
        fields = query.all()
        for f in fields:
            f_dict = {col.name: getattr(f, col.name) for col in f.__table__.columns}
            f_dict["_source_type"] = "field"
            results["fields"].append(f_dict)
            
    # 2. Search FinancialRecords
    if not type or type == "financial":
        query = db.query(FinancialRecord)
        if category:
            query = query.filter(FinancialRecord.category == category)
        if year:
            query = query.filter(FinancialRecord.fiscal_year == year)
        if q:
            search_str = f"%{q}%"
            query = query.filter(
                or_(
                    FinancialRecord.metric_label.ilike(search_str),
                    FinancialRecord.value.ilike(search_str)
                )
            )
        financials = query.all()
        for f in financials:
            f_dict = {col.name: getattr(f, col.name) for col in f.__table__.columns}
            f_dict["_source_type"] = "financial"
            results["financials"].append(f_dict)
            
    # 3. Search ProjectReferences
    if not type or type == "project":
        query = db.query(ProjectReference)
        if q:
            search_str = f"%{q}%"
            query = query.filter(
                or_(
                    ProjectReference.project_name.ilike(search_str),
                    ProjectReference.client_name.ilike(search_str),
                    ProjectReference.location.ilike(search_str),
                    ProjectReference.consultant.ilike(search_str),
                    ProjectReference.pmc.ilike(search_str)
                )
            )
        projects = query.all()
        for p in projects:
            p_dict = {col.name: getattr(p, col.name) for col in p.__table__.columns}
            p_dict["_source_type"] = "project"
            results["projects"].append(p_dict)
            
    # Add counts
    results["counts"] = {
        "fields": len(results["fields"]),
        "financials": len(results["financials"]),
        "projects": len(results["projects"])
    }
    
    return results

@router.get("/financial-records")
def get_financial_records(
    category: Optional[str] = None,
    year: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(FinancialRecord)
    if category:
        query = query.filter(FinancialRecord.category == category)
    if year:
        query = query.filter(FinancialRecord.fiscal_year == year)
    return query.all()

@router.get("/project-references")
def get_project_references(
    search: Optional[str] = None,
    region: Optional[str] = None,
    status: Optional[str] = None,
    client: Optional[str] = None,
    pmc: Optional[str] = None,
    third_party: Optional[str] = None,
    sector: Optional[str] = None,
    min_value: Optional[float] = None,
    max_value: Optional[float] = None,
    min_area: Optional[float] = None,
    max_area: Optional[float] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db)
):
    query = db.query(ProjectReference)
    if region:
        query = query.filter(ProjectReference.region == region)
    if status:
        query = query.filter(ProjectReference.status == status)
    if client:
        query = query.filter(ProjectReference.client_name == client)
    if pmc:
        query = query.filter(ProjectReference.pmc == pmc)
    if third_party:
        query = query.filter(ProjectReference.third_party == third_party)
    if sector:
        query = query.filter(ProjectReference.project_sector == sector)
    if search:
        s = f"%{search}%"
        query = query.filter(or_(
            ProjectReference.project_name.ilike(s),
            ProjectReference.client_name.ilike(s),
            ProjectReference.location.ilike(s),
            ProjectReference.pmc.ilike(s),
            ProjectReference.third_party.ilike(s),
        ))
    if min_value is not None:
        query = query.filter(ProjectReference.project_value_cr >= min_value)
    if max_value is not None:
        query = query.filter(ProjectReference.project_value_cr <= max_value)
    if min_area is not None:
        query = query.filter(ProjectReference.area_sqft_numeric >= min_area)
    if max_area is not None:
        query = query.filter(ProjectReference.area_sqft_numeric <= max_area)

    total = query.count()
    results = query.order_by(ProjectReference.id.desc()).offset((page - 1) * page_size).limit(page_size).all()
    return {
        "items": [
            {col.name: getattr(p, col.name) for col in p.__table__.columns}
            for p in results
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@router.get("/project-references/filter-options")
def get_project_reference_filter_options(db: Session = Depends(get_db)):
    def distinct_values(col):
        return sorted(v[0] for v in db.query(col).filter(col.isnot(None), col != '').distinct().all())
    return {
        "regions": distinct_values(ProjectReference.region),
        "statuses": distinct_values(ProjectReference.status),
        "pmcs": distinct_values(ProjectReference.pmc),
        "third_parties": distinct_values(ProjectReference.third_party),
        "sectors": distinct_values(ProjectReference.project_sector),
    }


# Exact-name column mapping for the unified project registry export (built by
# consolidating the File Cabinet's "Project Registry" Excel files) -- unlike
# company_data.py's /import-projects (fuzzy header matching, upserts by exact
# project_name), this expects a known fixed header row and inserts every row
# unconditionally. That's deliberate: the source data legitimately contains
# multiple rows sharing the same project_name (the same project logged in more
# than one source file with different figures), and upserting by name would
# silently collapse and lose those rows.
_BULK_IMPORT_REQUIRED_HEADERS = [
    "Project Name", "Client Name", "Location", "Scope of Work / Work Type",
    "Area (sq.ft)", "Cost (Rs Cr)", "Client Segment", "Source File", "Notes",
]

_AREA_RE = re.compile(r'^\s*([\d,]+(?:\.\d+)?)\s*sq\.?\s*ft\.?\s*$', re.IGNORECASE)
_COST_RE = re.compile(r'^\s*Rs\.?\s*([\d,]+(?:\.\d+)?)\s*Cr\.?\s*$', re.IGNORECASE)


def _extract_notes_field(notes: str, key: str) -> Optional[str]:
    """Notes cells look like 'Status: Ongoing; PMC: CBRE; Client contact: ...'
    -- pull out one 'Key: value' segment (value runs to the next '; ' or end)."""
    if not notes:
        return None
    m = re.search(re.escape(key) + r'\s*:\s*(.*?)(?:;\s*[A-Za-z][A-Za-z \-/]*:|$)', notes)
    return m.group(1).strip() or None if m else None


@router.post("/project-references/bulk-import")
async def bulk_import_project_references(file: UploadFile = File(...), db: Session = Depends(get_db)):
    from utils import enforce_upload_size
    contents = await file.read()
    enforce_upload_size(len(contents))
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
    sheet_name = "Project Registry" if "Project Registry" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h else "" for h in next(rows_iter)]
    missing = [h for h in _BULK_IMPORT_REQUIRED_HEADERS if h not in headers]
    if missing:
        wb.close()
        raise HTTPException(400, f"Uploaded file is missing expected column(s): {', '.join(missing)}")
    col_idx = {h: i for i, h in enumerate(headers)}

    imported = 0
    skipped_blank_rows = 0
    rows_in_sheet = 0
    for row in rows_iter:
        rows_in_sheet += 1
        if not any(row):
            skipped_blank_rows += 1
            continue

        def get(h):
            v = row[col_idx[h]] if col_idx[h] < len(row) else None
            return str(v).strip() if v is not None else ""

        area_raw = get("Area (sq.ft)")
        cost_raw = get("Cost (Rs Cr)")
        notes_raw = get("Notes")

        area_m = _AREA_RE.match(area_raw)
        cost_m = _COST_RE.match(cost_raw)

        row_data = {
            "project_name": get("Project Name") or None,
            "client_name": get("Client Name") or None,
            "location": get("Location") or None,
            "project_type": get("Scope of Work / Work Type") or None,
            "project_sector": get("Client Segment") or None,
            "source_file": get("Source File") or file.filename,
            "area_sqft": area_raw or None,
            "area_sqft_numeric": float(area_m.group(1).replace(",", "")) if area_m else None,
            "project_value": cost_raw or None,
            "project_value_cr": float(cost_m.group(1).replace(",", "")) if cost_m else None,
            "notes": notes_raw or None,
            "pmc": _extract_notes_field(notes_raw, "PMC"),
            "third_party": _extract_notes_field(notes_raw, "Third party"),
            "consultant": _extract_notes_field(notes_raw, "Consultant"),
            "status": _extract_notes_field(notes_raw, "Status"),
            "start_date": _extract_notes_field(notes_raw, "Start"),
            "end_date": _extract_notes_field(notes_raw, "End"),
        }
        db.add(ProjectReference(**row_data))
        imported += 1

    db.commit()
    wb.close()
    return {"imported": imported, "rows_in_sheet": rows_in_sheet, "skipped_blank_rows": skipped_blank_rows}
